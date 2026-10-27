"""A4 print renderers for 学案 and 试卷 (python-docx).

Typography comes from the teacher's own material (25 学案 / 147 试卷 in the corpus):
  * A4, margins 1.9cm; 标题 18pt 华文中宋居中; 栏目 13pt 黑体居中 + 分隔线
  * 正文 宋体 12pt (小四), 行距 1.5; 项目符号用悬挂缩进
  * 填空用下划线空格（宽度随答案字数），答案/解析用浅底纹
  * 页脚：学校 + 第 X 页 共 Y 页
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .mdparse import Block, Doc, blanks, parse

CN_BODY = "宋体"
CN_HEAD = "黑体"
CN_TITLE = "华文中宋"
EN_BODY = "Times New Roman"
GREY = RGBColor(0x59, 0x59, 0x59)
ANSWER_INK = RGBColor(0x8A, 0x1B, 0x1B)

PAGE_W, PAGE_H = Cm(21.0), Cm(29.7)


# ------------------------------------------------------------------ helpers
def _style_run(run, cn=CN_BODY, en=EN_BODY, size=12, bold=False, color=None, underline=False):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.underline = underline
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cn)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)


def _para(doc, text="", *, cn=CN_BODY, en=EN_BODY, size=12, bold=False,
          align=None, first_indent=None, left_indent=None, hanging=None,
          space_before=0, space_after=0, line=1.5, color=None, style=None,
          keep_next=False):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    if left_indent is not None:
        pf.left_indent = Cm(left_indent)
    if hanging is not None:
        pf.first_line_indent = Cm(-hanging)
    elif first_indent is not None:
        pf.first_line_indent = Cm(first_indent)
    if align is not None:
        p.alignment = align
    if text:
        _style_run(p.add_run(text), cn, en, size, bold, color)
    return p


def _blank_run(p, chars: int, size=12):
    # 全角空格，1 个 ≈ 1 个汉字宽，下划线连续不断
    r = p.add_run("\u3000" * max(1, chars))
    _style_run(r, CN_BODY, EN_BODY, size, underline=True)


LABEL_RE = re.compile(r"^(重点|难点|考点|例\s*\d+|探究[一二三四五六七八九十]?|变式[一二三四五六七八九十]?|思考|提示|归纳|答案|解析)\s*[:：]")


def _write_runs(p, text: str, size: float, color=None, bold=False):
    """Write a paragraph body, turning {{答案}} into underlined blanks."""
    m = LABEL_RE.match(text)
    if m:
        _style_run(p.add_run(m.group(0)), CN_HEAD, EN_BODY, size, True, color)
        text = text[m.end():]
    for kind, val in blanks(text):
        if kind == "blank":
            _blank_run(p, max(4, min(14, val + 2)), size)
        else:
            _style_run(p.add_run(val), CN_BODY, EN_BODY, size, bold, color)


def _field(p, instr: str, size=9):
    run = p.add_run()
    _style_run(run, CN_BODY, EN_BODY, size, color=GREY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(it)
    run._r.append(end)


def _footer(section, school: str = ""):
    footer = section.footer
    first = footer.paragraphs[0]
    if school:
        first.alignment = WD_ALIGN_PARAGRAPH.LEFT
        first.paragraph_format.space_after = Pt(0)
        _style_run(first.add_run(school), CN_BODY, EN_BODY, 9, color=GREY)
    p = footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    _style_run(p.add_run("第 "), CN_BODY, EN_BODY, 9, color=GREY)
    _field(p, "PAGE", 9)
    _style_run(p.add_run(" 页　共 "), CN_BODY, EN_BODY, 9, color=GREY)
    _field(p, "NUMPAGES", 9)
    _style_run(p.add_run(" 页"), CN_BODY, EN_BODY, 9, color=GREY)


def _table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "999999")
        borders.append(el)
    tblPr.insert_element_before(
        borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
        "w:tblCaption", "w:tblDescription", "w:tblPrChange")


def _shade(p, fill="F4F1F7"):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def _rule(p, color="0B4FA8"):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    pPr.append(pbdr)


def new_doc(margins=(Cm(1.9), Cm(1.9), Cm(2.0), Cm(2.0))):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = PAGE_W, PAGE_H
    s.left_margin, s.right_margin = margins[0], margins[1]
    s.top_margin, s.bottom_margin = margins[2], margins[3]
    normal = doc.styles["Normal"]
    normal.font.name = EN_BODY
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN_BODY)
    for name, size, cn in (("Heading 1", 18, CN_TITLE), ("Heading 2", 13, CN_HEAD),
                           ("Heading 3", 12, CN_HEAD)):
        st = doc.styles[name]
        st.font.name = EN_BODY
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        rf.set(qn("w:eastAsia"), cn)
    return doc


def _heading(doc, text, level, *, center=False, size=None, rule=False):
    p = doc.add_paragraph(style=f"Heading {level}")
    pf = p.paragraph_format
    pf.space_before = Pt(12 if level == 2 else 8)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.2
    pf.keep_with_next = True
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(p.add_run(text), CN_HEAD if level > 1 else CN_TITLE,
               EN_BODY, size or (13 if level == 2 else 12), bold=True)
    if rule:
        _rule(p)
    return p


# ------------------------------------------------------------------ blocks
def _render_blocks(doc, blocks, *, body_size=12, indent=0.74, line=1.5):
    for b in blocks:
        if b.kind == "text":
            p = _para(doc, "", size=body_size, first_indent=indent, line=line,
                      space_after=2)
            _write_runs(p, b.text, body_size)
        elif b.kind == "note":
            _para(doc, b.text, size=body_size - 1.5, color=GREY, first_indent=indent,
                  line=1.3, space_after=2)
        elif b.kind in ("bullet", "num"):
            p = _para(doc, "", size=body_size, left_indent=indent + 0.6 * b.level,
                      hanging=indent, line=line, space_after=3)
            if b.kind == "bullet":
                _style_run(p.add_run("▪ "), CN_BODY, EN_BODY, body_size - 1)
            _write_runs(p, b.text, body_size)
        elif b.kind == "table":
            if not b.rows:
                continue
            cols = max(len(r) for r in b.rows)
            t = doc.add_table(rows=len(b.rows), cols=cols)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            _table_borders(t)
            for ri, row in enumerate(b.rows):
                for ci in range(cols):
                    cell = t.cell(ri, ci)
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.line_spacing = 1.15
                    _style_run(p.add_run(row[ci] if ci < len(row) else ""),
                               CN_BODY, EN_BODY, body_size - 1.5, bold=(ri == 0))
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif b.kind == "image":
            path = Path(b.src)
            if path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(path), width=Cm(9))
                if b.text:
                    _para(doc, b.text, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY)
        elif b.kind == "answer":
            p = _para(doc, "", size=body_size, left_indent=0.4, line=line, space_after=2)
            _shade(p)
            _write_runs(p, b.text, body_size, color=ANSWER_INK)
        elif b.kind == "break":
            doc.add_page_break()


# ------------------------------------------------------------------ 学案
def render_xuean(md: str, out_path: str | Path, *, school: str = "") -> Path:
    d = parse(md)
    doc = new_doc()
    _footer(doc.sections[0], school)

    _para(doc, d.title or "学案", cn=CN_TITLE, size=18, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, line=1.15)
    meta = "｜".join(x for x in [school, d.meta.get("学科"), d.meta.get("年级"),
                                 f"课时：{d.meta['课时']}" if d.meta.get("课时") else "",
                                 d.meta.get("编写")] if x)
    if meta:
        _para(doc, meta, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY,
              space_after=8, line=1.1)

    if d.preface:
        _render_blocks(doc, d.preface)

    for s in d.sections:
        if not s.title:
            _render_blocks(doc, s.blocks)
            continue
        if s.level == 2:
            _heading(doc, s.title, 2, center=True, rule=True)
        else:
            _heading(doc, s.title, 3)
        _render_blocks(doc, s.blocks)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path



def _seal_line(doc):
    """班级 / 姓名 / 考号 三栏，用表格底框线做填写线（中间留空隙，三段独立）。"""
    usable = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    widths = [0.30, 0.05, 0.26, 0.05, 0.34]
    labels = {0: "班级：", 2: "姓名：", 4: "考号："}
    table = doc.add_table(rows=1, cols=5)
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for ci, ratio in enumerate(widths):
        cell = table.cell(0, ci)
        cell.width = int(usable * ratio)
        para = cell.paragraphs[0]
        para.paragraph_format.space_before = Pt(2)
        para.paragraph_format.space_after = Pt(6)
        if ci not in labels:
            continue
        tcPr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "808080")
        borders.append(bottom)
        tcPr.append(borders)
        _style_run(para.add_run(labels[ci]), CN_BODY, EN_BODY, 10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def _expand_options(blocks: list[Block]) -> list[Block]:
    """把 'A.甲　B.乙　C.丙　D.丁' 这种一行多选项拆成独立块。"""
    out = []
    for b in blocks:
        if b.kind in ("text", "bullet") and b.text:
            t = b.text.strip()
            marks = list(INLINE_OPT.finditer(t))
            if len(marks) >= 3:
                for i, m in enumerate(marks):
                    end = marks[i + 1].start() if i + 1 < len(marks) else len(t)
                    piece = t[m.start():end].strip()
                    if piece:
                        out.append(Block(b.kind, text=piece))
                continue
        out.append(b)
    return out

# ------------------------------------------------------------------ 试卷
OPTION_RE = re.compile(r"^([A-DＡ-Ｄ])\s*[.．、]\s*(.*)$")
INLINE_OPT = re.compile(r"([A-DＡ-Ｄ])\s*[.．、]\s*")


def render_paper(md: str, out_path: str | Path, *, school: str = "",
                 answer_page: bool = True, seal_line: bool = True) -> Path:
    d = parse(md)
    doc = new_doc(margins=(Cm(2.0), Cm(2.0), Cm(2.2), Cm(2.0)))
    _footer(doc.sections[0], school)

    _para(doc, d.title or "试卷", cn=CN_TITLE, size=18, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, line=1.15)
    meta_bits = [f"{k}：{d.meta[k]}" for k in ("时间", "满分") if d.meta.get(k)]
    if meta_bits:
        _para(doc, "（" + "　".join(meta_bits) + "）", size=10.5,
              align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY, space_after=4)
    if seal_line:
        _seal_line(doc)

    in_answers = False
    for s in d.sections:
        if re.match(r"^(参考答案|答案与解析|答案及解析|答案|解析)", s.title.strip()):
            in_answers = True
            if answer_page:
                doc.add_page_break()
            _para(doc, s.title, cn=CN_HEAD, size=14, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
            _render_blocks(doc, s.blocks, body_size=11)
            continue
        if in_answers:
            _render_blocks(doc, s.blocks, body_size=11)
            continue
        if s.title:
            _heading(doc, s.title, 2, center=False, size=12)
        blocks = _expand_options(s.blocks)
        i = 0
        while i < len(blocks):
            b = blocks[i]
            if b.kind in ("text", "bullet") and OPTION_RE.match(b.text.strip()):
                row = []
                while i < len(blocks) and blocks[i].kind in ("text", "bullet") \
                        and OPTION_RE.match(blocks[i].text.strip()):
                    row.append(blocks[i])
                    i += 1
                _option_row(doc, row)
                continue
            _paper_block(doc, b)
            i += 1

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def _tidy(text: str) -> str:
    """避免 '（　　）' 孤行：括号内用不换行空格，并与前一个汉字绑定。"""
    text = re.sub(r"（[\s\u3000]*）", "（\u00a0\u00a0）", text)
    text = re.sub(r"(?<=[\u4e00-\u9fa5])\s*（\u00a0", "\u00a0（\u00a0", text)
    return text


def _option_row(doc, row: list[Block]):
    """同行 4 项用无框表格排，宽度不够就 2 项一行，保证竖向对齐、不断词。"""
    def cw(t: str) -> float:
        return sum(1.0 if ord(c) > 127 else 0.55 for c in t)

    total = sum(cw(b.text.strip()) for b in row) + 2.0 * (len(row) - 1)
    per = 4 if total <= 54 and len(row) == 4 else (2 if len(row) >= 2 else 1)
    if len(row) <= 2:
        per = 1
    usable = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    rows = [row[i:i + per] for i in range(0, len(row), per)]
    table = doc.add_table(rows=len(rows), cols=per)
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for ri, r in enumerate(rows):
        for ci in range(per):
            cell = table.cell(ri, ci)
            cell.width = int(usable / per)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.line_spacing = 1.3
            if ci < len(r):
                _style_run(para.add_run(r[ci].text.strip()), CN_BODY, EN_BODY, 11.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def _paper_block(doc, b: Block, keep_next: bool = False):
    if b.kind == "text":
        text = _tidy(b.text.strip())
        if re.match(r"^\d{1,2}[.．、]", text):
            p = _para(doc, "", size=12, left_indent=0.74, hanging=0.74, line=1.5,
                      space_after=3, keep_next=keep_next)
        elif OPTION_RE.match(text):
            p = _para(doc, "", size=11.5, left_indent=0.74, hanging=0.74, line=1.35,
                      space_after=2, keep_next=keep_next)
        elif text.startswith("（") and text.endswith("）") and len(text) < 40:
            _para(doc, text, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
            return
        else:
            p = _para(doc, "", size=12, first_indent=0.74, line=1.5, space_after=3)
        _write_runs(p, text, 12 if not OPTION_RE.match(text) else 11.5)
    elif b.kind == "num":
        p = _para(doc, "", size=12, left_indent=0.74, hanging=0.74, line=1.5,
                  space_after=3, keep_next=keep_next)
        _write_runs(p, b.text, 12)
    elif b.kind == "answer":
        p = _para(doc, "", size=11.5, left_indent=0.4, line=1.5, space_after=2)
        _shade(p)
        _write_runs(p, b.text, 11.5, color=ANSWER_INK)
    elif b.kind == "bullet":
        p = _para(doc, "", size=11.5, left_indent=0.74, hanging=0.74, line=1.45, space_after=2)
        _style_run(p.add_run("▪ "), CN_BODY, EN_BODY, 10.5)
        _write_runs(p, b.text, 11.5)
    elif b.kind in ("table", "image"):
        _render_blocks(doc, [b], body_size=11.5)
    elif b.kind == "break":
        doc.add_page_break()
    else:
        _render_blocks(doc, [b], body_size=11)
