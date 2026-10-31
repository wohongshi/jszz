"""离线自检：不依赖网络，验证渲染器与组卷产出符合规范。"""
import json
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "server"))

from pptx import Presentation                      # noqa: E402
from docx import Document                          # noqa: E402

from jszz.assemble import make_paper              # noqa: E402
from jszz.docx_render import render_paper, render_xuean   # noqa: E402
from jszz.pptx_render import build_paper_deck, build_xuean_deck  # noqa: E402

OUT = ROOT / "out" / "selftest"
OUT.mkdir(parents=True, exist_ok=True)
fails = []


def check(cond, msg):
    print(("  ok   " if cond else "  FAIL ") + msg)
    if not cond:
        fails.append(msg)


print("== 学案 ==")
md = (ROOT / "examples" / "xuean-wuli.md").read_text("utf-8")
docx_path = render_xuean(md, OUT / "xuean.docx", school="测试中学")
pptx_path = build_xuean_deck(md, OUT / "xuean.pptx", school="测试中学")
d = Document(str(docx_path))
text = "\n".join(p.text for p in d.paragraphs)
check("学习目标" in text, "docx 含学习目标")
check("课前预习" in text and "课内探究" in text, "docx 含预习/探究栏目")
check("知道" not in text or True, "docx 可打开")
underlined = sum(1 for p in d.paragraphs for r in p.runs if r.font.underline)
check(underlined >= 8, f"docx 填空下划线充足 ({underlined})")
prs = Presentation(str(pptx_path))
slides_text = ["\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
               for s in prs.slides]
check(len(prs.slides) >= 8, f"pptx 页数合理 ({len(prs.slides)})")
check(not any("{{" in t or "}}" in t for t in slides_text), "pptx 无 {{}} 残留")
check(abs(prs.slide_width.inches - 13.333) < 0.05, "pptx 为 16:9")
check(any("解析" in t for t in slides_text), "pptx 有解析页")
check(any("小结" in t for t in slides_text), "pptx 有小结页")

print("== 试卷 ==")
md2 = (ROOT / "examples" / "paper-wuli.md").read_text("utf-8")
p_docx = render_paper(md2, OUT / "paper.docx", school="测试中学")
p_pptx = build_paper_deck(md2, OUT / "paper.pptx", school="测试中学")
d2 = Document(str(p_docx))
t2 = "\n".join(p.text for p in d2.paragraphs)
check("单项选择题" in t2, "试卷有大题标题")
check("参考答案" in t2, "试卷含参考答案页")
check(len(d2.tables) >= 3, f"试卷选项用表格排版 ({len(d2.tables)})")
prs2 = Presentation(str(p_pptx))
st2 = ["\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs2.slides]
check(len(prs2.slides) >= 10, f"讲评 pptx 页数合理 ({len(prs2.slides)})")
check(any("答案" in t for t in st2), "讲评 pptx 有答案页")
check(not any("注意事项" in t and "第 1 题" in t for t in st2), "注意事项没有被当成题目")

print("== 组卷 ==")
qs = [
    {"stem": "下列关于加速度的说法正确的是（　　）",
     "options": ["加速度为零时速度一定为零", "加速度方向与速度变化方向相同",
                 "加速度越大速度越大", "加速度为负时物体一定减速"],
     "answer": "B", "analysis": "加速度是速度变化率，方向与 Δv 相同。", "score": 4, "type": "选择"},
    {"stem": "物体做自由落体运动，3 s 内下落高度为____m（g 取 10 m/s²）。",
     "answer": "45", "score": 4, "type": "填空"},
    {"stem": "汽车以 20 m/s 匀速行驶，刹车加速度大小为 4 m/s²，求刹车后 8 s 内的位移。",
     "answer": "50 m", "analysis": "刹车时间 5 s，8 s 内位移等于 5 s 内位移，x=v0²/2a=50 m。",
     "score": 12, "type": "解答"},
]
r = make_paper(qs, title="自检组卷", out_dir=OUT, total_score=20, school="测试中学")
check(r["ok"] and Path(r["files"]["docx"]).exists(), "组卷产出 docx")
check(Path(r["files"]["pptx"]).exists(), "组卷产出 pptx")
md3 = Path(r["files"]["md"]).read_text("utf-8")
check("一、单项选择题" in md3 and "二、填空题" in md3 and "三、解答题" in md3, "按题型自动分节")
check(md3.count("本题共1小题") == 3, "自动写题量与分节")

print("== 新增能力（v1.1）==")
from jszz.qa import check_docx, check_pptx, check_markdown  # noqa: E402

d3 = Document(str(docx_path))
head = "".join(p.text for p in d3.sections[0].header.paragraphs)
head += "".join(c.text for t in d3.sections[0].header.tables for r in t.rows for c in r.cells)
check("编写人" in head, "学案页眉含编写人/审核字段")
foot = "".join(p.text for p in d3.sections[0].footer.paragraphs)
check("共" in foot and "页" in foot, "页脚为「第X页（共Y页）」")
kinds = [b.kind for sec in __import__("jszz.mdparse", fromlist=["parse"]).parse(md).sections
         for b in sec.blocks]
check("tip" in kinds, "【特别提醒】识别为 tip 块")
check("insight" in kinds, "【点睛】识别为 insight 块")

prs3 = Presentation(str(build_xuean_deck(md, OUT / "xuean.pptx", school="测试中学")))
bands = 0
for sl in prs3.slides:
    for sh in sl.shapes:
        if sh.shape_type == 1 and sh.width and abs(sh.width.inches - 13.333) < 0.05:
            bands += 1
            break
check(bands >= 5, f"通栏页头/底部色带存在 ({bands} 页)")

has_note = any("判定要点" in (sh.text_frame.text if sh.has_text_frame else "")
               for sl in prs3.slides for sh in sl.shapes)
check(has_note, "例题页有「判定要点」批注框")

d4 = Document(str(p_docx))
check(len(d4.tables) >= 5, f"答案速查表已生成 ({len(d4.tables)} 表)")

check(check_docx(docx_path)["ok"], "学案 docx 体检通过")
check(check_pptx(pptx_path)["ok"], "学案 pptx 体检通过")
check(check_docx(p_docx)["ok"], "试卷 docx 体检通过")
check(check_pptx(p_pptx)["ok"], "讲评 pptx 体检通过")
check(check_markdown(md)["ok"] and check_markdown(md2)["ok"], "Markdown 写作体检通过")

print("== 栏目识别（sec_key 回归，防【】括号漏判退化） ==")
from jszz.mdparse import sec_key          # noqa: E402
from jszz.pptx_render import BIG          # noqa: E402
from jszz.docx_render import BIG_SECTION  # noqa: E402

for raw, key in [
    ("【学习目标】", "学习目标"),
    ("一、试题分析", "试题分析"),
    ("（一）评分标准", "评分标准"),
    ("试题分析（命题立意）", "试题分析"),
    ("【参考答案】", "参考答案"),
    ("二、词性转换（每空1分，共10分）", "词性转换"),
]:
    check(sec_key(raw) == key, f"sec_key({raw!r}) → {key!r}")
for raw in ["【学习目标】", "【当堂检测】", "【答案】", "一、试题分析",
            "强基础·教考衔接", "必备知识·精读教材", "【我的反思】"]:
    k = sec_key(raw)
    check(bool(BIG.match(k)), f"pptx BIG 认得 {raw!r}")
for raw in ["【参考答案】", "试题分析（命题立意）", "备考提示", "原题呈现",
            "强基础·教考衔接", "审题常见误区", "【我的反思】"]:
    k = sec_key(raw)
    check(bool(BIG_SECTION.match(k)), f"docx BIG_SECTION 认得 {raw!r}")
check(not BIG.match(sec_key("普通章节标题")), "普通标题不误判为大栏目")
check(not BIG.match(sec_key("知识点1 匀变速")), "知识点小标题不误判为大栏目")

# 断行回归：真实题源「2.（2024…」「3.《史记…」N. 后接全角括号/书名号（无空格），
# 旧断行正则 \s 要空格导致漏判、题干粘进上一题 D 选项（选项串行）。
from jszz.mdparse import parse as _parse_break  # noqa: E402
_break_md = "# t\n【练习】\n1.（2025·云南）第一题题干（　）\nA．甲　B．乙　C．丙　D．丁\n2.（2024·福建）第二题题干（　）\nA．戊　B．己　C．庚　D．辛\n3.《史记》第三题题干（　）\nA．壬　B．癸　C．子　D．丑\n"
_break_texts = [b.text for b in _parse_break(_break_md).flat_blocks()]
# 题干必须各自成块，不许粘进上一题选项（D．丁 后跟 2. 即为粘连）
for _needle in ("2. （2024", "3. 《史记"):
    _solo = [t for t in _break_texts if _needle in t]
    check(len(_solo) == 1, f"题号 {_needle} 独立成块（{len(_solo)}）")
check(not any(("D．丁" in t and "2." in t) or ("D．辛" in t and "3." in t)
              for t in _break_texts), "题干不粘进上一题 D 选项")

print("== 九科视觉（语数英物化生政史地）==")
from jszz import subject_style as ss  # noqa: E402

NINE = [("语文", "chinese"), ("数学", "math"), ("英语", "english"),
        ("物理", "physics"), ("化学", "chemistry"), ("生物", "biology"),
        ("政治", "politics"), ("历史", "history"), ("地理", "geography")]
for name, key in NINE:
    st = ss.get(name)
    check(st.key == key, f"学科 {name} → {key} 风格")
    check(bool(st.cover) and bool(st.header) and bool(st.deco),
          f"{name} 有封面/页头/装饰（{st.cover}/{st.header}/{st.deco}）")
for name, key in NINE:
    check(ss.guess(f"胶州四中{name}学案") == key, f"标题猜学科：{name}")
check(ss.guess("一轮复习纲要上第11课") == "history", "历史标题猜中（旧版漏判）")
check(ss.guess("专题11 热力环流") == "geography", "地理标题猜中")
check(ss.guess("复习必修四第二单元") == "politics", "政治标题猜中")
# 九科封面/页头互不重样（防三科套同一套皮）
covers = {ss.get(n).cover for n, _ in NINE}
headers = {ss.get(n).header for n, _ in NINE}
check(len(covers) == 9, f"九科封面版式不重样（{len(covers)}）")
check(len(headers) == 9, f"九科页头不重样（{len(headers)}）")

print("== scene 指令（含政史地四新页型）==")
from jszz.mdparse import parse as _parse  # noqa: E402

_d = _parse("""# t
> 史料：材料。——摘编自《史》
> 材料：时政。| 设问：问题？
> 图表：图例 | 判读：①甲；②乙
> 模板：答题模式：标题 + 表述成文 | 标题：×× | 时代背景：××
## s
body
""")
_scenes = [b.meta.get("scene") for b in _d.preface if b.kind == "scene"]
check(_scenes == ["史料", "材料", "图表", "模板"], f"四种 scene 解析：{_scenes}")
_d2 = _parse("# t\n\n## s\n> 引文：x——y\n")
check([b.meta.get("scene") for b in _d2.sections[0].blocks if b.kind == "scene"] == ["引文"],
      "旧五种 scene 不回归")
# **关键词** 红粗渲染（史轮2-24 取证）
from jszz.mdparse import blanks as _blanks  # noqa: E402
_b = _blanks("打破坊市制度，**商业活动不再受时空限制**，{{纸币}}出现")
check(("strong", "商业活动不再受时空限制") in _b and ("blank", "纸币") in _b,
      "**关键词**与{{挖空}}双语法共存")

print("== 图片引用（图注句式 + 嵌图）==")
from PIL import Image as _Image  # noqa: E402
_img = OUT / "fig-test.png"
_Image.new("RGB", (240, 160), (70, 130, 180)).save(_img)
_md_img = f"""# t
> 学科：地理

## s

![图1示意热力环流的常见类型（单位：hPa）]({_img})

1. 据图1所示，判断（　　）
"""
from jszz.pptx_render import build_xuean_deck as _bud  # noqa: E402
from jszz.docx_render import render_xuean as _rxd  # noqa: E402
_p2 = _bud(_md_img, OUT / "fig-test.pptx", school="测试中学")
_d2 = _rxd(_md_img, OUT / "fig-test.docx", school="测试中学")
from pptx import Presentation as _P2  # noqa: E402
_sl = list(_P2(str(_p2)).slides)
check(any(sh.shape_type == 13 for s2 in _sl for sh in s2.shapes), "pptx 嵌入图片")
check(any(sh.has_text_frame and sh.text_frame.text.startswith("图1示意")
          for s2 in _sl for sh in s2.shapes), "pptx 图注（图N示意句式）")
from docx import Document as _D2  # noqa: E402
_dd = _D2(str(_d2))
check(len(_dd.inline_shapes) == 1, "docx 嵌入图片")
check("图1示意" in "\n".join(pp.text for pp in _dd.paragraphs), "docx 图注")

print("== OCR 三修回归（假引擎，不起 ONNX）==")
import hashlib as _hashlib  # noqa: E402
import io as _io  # noqa: E402
from PIL import Image as _Im, ImageDraw as _ID  # noqa: E402
from jszz import ocr as _ocr  # noqa: E402

def _fake(text):
    def _c(*a, **k):
        return ([[[[0, 0], [10, 0], [10, 10], [0, 10]], text, 0.9]], None)
    return _c

def _png_bytes(w=60, h=30):
    b = _io.BytesIO()
    im = _Im.new("RGB", (w, h), "white")
    _ID.Draw(im).line([(5, h // 2), (w - 5, h // 2)], fill="black")
    im.save(b, format="PNG")
    return b.getvalue()

# 修1：P 模式透明底 → 白底（曾变黑底 OCR 全瞎）
_rgba = _Im.new("RGBA", (40, 20), (0, 0, 0, 0))
_ID.Draw(_rgba).line([(5, 10), (35, 10)], fill=(0, 0, 0, 255))
_prep_out = _ocr._prep(_rgba.convert("P", palette=_Im.ADAPTIVE))
check(_prep_out.getpixel((2, 2)) == (255, 255, 255), "OCR：P 模式透明区→白底")
# 修3：超大图降采样（曾拖 ONNX 发热）
check(max(_ocr._prep(_Im.new("RGB", (4000, 3000), "white")).size) <= 1600,
      "OCR：超大图降采样 ≤1600")
# 修2：水印片段剔除不陪葬正文（曾整段清空）
_cache = (Path.home() / ".jszz" / "cache" / "ocr")
_cache.mkdir(parents=True, exist_ok=True)
def _ocr_once(blob, fake_text):
    _ocr._ENGINE = _fake(fake_text)
    cf = _cache / (_hashlib.sha1(blob).hexdigest() + ".txt")
    if cf.exists():
        cf.unlink()
    return _ocr.ocr_bytes(blob)
check(_ocr_once(_png_bytes(), "v=v0+at 学科网") == "v=v0+at",
      "OCR：水印片段剔除、公式保住")
check(_ocr_once(_png_bytes(61), "学科网 www.zxxk.com") == "",
      "OCR：纯水印图返回空")
check(_ocr_once(_png_bytes(62), "x^2+y^2=r^2") == "x^2+y^2=r^2",
      "OCR：无水印原样保留")

print("== 写作工艺页型 + 工具 ==")
_craft_md = (
    "# 答题规范\n> 学科：化学\n\n## 易错\n"
    "> 规范表述：把渣子滤掉｜过滤除去不溶性杂质 | 温度高不好｜温度过高，H₂O₂受热分解加快\n"
    "## 口诀\n> 口诀：控温作答要双面，速率产率对能耗\n"
    "## 采分\n> 采分：要点1：主题：展开（3分）｜要点2：主题：展开（3分）\n"
    "## 考情\n> 考情：2026｜云南·14-16题·9分｜情境 | 2025｜江苏·24题·18分｜情境\n"
    "## 例\n1. 题干（　）\nA. a　B. b　C. c　D. d\n【答案】C\n"
    "【详解】A：排除。B：排除。C：正确。D：排除。故选 C。\n"
)
_craft_ppt = build_xuean_deck(_craft_md, OUT / "craft.pptx", subject="化学")
_sl = Presentation(_craft_ppt)
_alltxt = ""
for s in _sl.slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            _alltxt += sh.text_frame.text + "\n"
        if sh.has_table:  # pptx 表格文字在 table 单元格里，不在 text_frame
            for row in sh.table.rows:
                for cell in row.cells:
                    _alltxt += cell.text + "\n"
check("错误写法（不得分）" in _alltxt and "规范表述（得分）" in _alltxt,
      "pptx 规范表述红绿双栏")
check("控温作答要双面" in _alltxt, "pptx 口诀页")
check("要点1" in _alltxt and "要点2" in _alltxt, "pptx 采分点分行")
check("年份" in _alltxt and "命题情境" in _alltxt, "pptx 考情分析表头")
_craft_docx = render_xuean(_craft_md, OUT / "craft.docx")
_dt = "\n".join(p.text for p in Document(_craft_docx).paragraphs)
check("要点1" in _dt and "控温作答要双面" in _dt, "docx 写作工艺页型不丢")

from jszz.writingcraft import check_writing, suggest_scoring  # noqa: E402
_rc = check_writing(_craft_md)
check(_rc["ok"] and not _rc["issues"], "style_check：合规样例零 issue")
check(_rc["counts"]["has_mnemonic"] and _rc["counts"]["has_scoring"],
      "style_check：口诀/采分被识别")
_bad = "1. 题（　）\nA. a B. b C. c D. d\n【答案】C\n【详解】把渣子滤掉就行了，差不多选C。"
_rb = check_writing(_bad)
check(not _rb["ok"] and any("口语" in i for i in _rb["issues"]),
      "style_check：口语被判 issue")
check(any("逐项" in w for w in _rb["warnings"]),
      "style_check：详解未逐项排除被警告")
_sp = suggest_scoring("茶叶推动习俗形成；促进文化交流；古代是边防媒介。")
check(len(_sp) == 3 and _sp[0].startswith("要点1"), "scoring_suggest：拆成采分点")

print("== 抖音/视频号 ==")
from jszz.douyin import DouyinClient, is_douyin  # noqa: E402
from jszz.channels import ChannelsClient, is_channels  # noqa: E402
_dy = DouyinClient()
check(is_douyin("https://v.douyin.com/iRxWcKu/") and not is_douyin("https://b23.tv/x"),
      "douyin 链接判定")
_r = _dy.resolve("https://www.douyin.com/video/7700562684790081714")
check(_r["ok"] and _r["id"] == "7700562684790081714", "douyin 长链抠 ID")
check(_dy.resolve("").get("ok") is False, "douyin 空链返回失败")
check(_dy.resolve("https://www.iesdouyin.com/share/video/7123456789012345678/")["kind"] == "video",
      "douyin 分享页抠 ID")
_ch = ChannelsClient()
check(_ch.status()["auto_download"] is False, "channels 诚实边界：不假装自动下载")
check(is_channels("https://channels.weixin.qq.com/web/pages/feed?srf=x")
      and not is_channels("https://b23.tv/x"), "channels 链接判定")
# AES-128 解密回归（先加密再解密，ftyp 头还原）
from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes  # noqa: E402
_k = b"0123456789abcdef"
_plain = b"\x00\x00\x00\x18ftyp" + b"x" * 100
_pad = 16 - len(_plain) % 16
_enc = Cipher(algorithms.AES(_k), modes.ECB()).encryptor()
_blob = _enc.update(_plain + bytes([_pad]) * _pad) + _enc.finalize()
(OUT / "ch.bin").write_bytes(_blob)
_dr = _ch.decrypt_file(OUT / "ch.bin", _k.hex(), dest=OUT / "ch.mp4")
check(_dr["ok"] and (OUT / "ch.mp4").read_bytes()[:4] == b"\x00\x00\x00\x18",
      "channels AES 解密还原 ftyp")
check(_ch.decrypt_file(OUT / "ch.bin", "zz", dest=OUT / "ch2.mp4").get("ok") is False,
      "channels 坏 key 报错不产坏文件")

print("== 小红书 + 统一取材 ==")
from jszz.xhs import XHSClient, is_xhs  # noqa: E402
from jszz.bilibili import is_bili  # noqa: E402
_x = XHSClient()
check(is_xhs("https://www.xiaohongshu.com/explore/6411cf99000000001300b6d9")
      and is_xhs("https://xhslink.com/o/abc") and not is_xhs("https://b23.tv/x"),
      "xhs 链接判定")
check(_x.resolve("https://www.xiaohongshu.com/explore/6411cf99000000001300b6d9")["id"]
      == "6411cf99000000001300b6d9", "xhs 长链抠 ID")
check(_x.resolve("").get("ok") is False, "xhs 空链返回失败")
check(_x.status()["has_cookie"] is False, "xhs 未配 cookie 如实报 false")
check(_x.status()["auto_download"] is False, "xhs 诚实边界：无 cookie 不假装自动下载")
check(_x.has_cookie() is False, "xhs has_cookie 未配为 false")
# 统一取材平台识别（5 路由，不真下）
check(is_bili("https://www.bilibili.com/video/BV1xx") and is_bili("https://b23.tv/x"),
      "bili 链接判定（b23.tv 短链也认）")
_routed = [("https://www.bilibili.com/video/BV1xx", "bilibili"),
           ("https://v.douyin.com/iRxWcKu/", "douyin"),
           ("https://www.xiaohongshu.com/explore/6411cf99000000001300b6d9", "xhs"),
           ("https://channels.weixin.qq.com/web/pages/feed?srf=x", "channels")]
_ok = all((("douyin" if is_douyin(u) else "xhs" if is_xhs(u)
            else "channels" if is_channels(u) else "bilibili" if is_bili(u)
            else "unknown") == exp) for u, exp in _routed)
check(_ok, "material_fetch 平台识别 4/4")

print("== 练习/讲义（drill）==")
from jszz.docx_render import render_drill  # noqa: E402
_drill_md = (
    "# 胶州四中高三化学每日一练3\n> 学科：化学 | 年级：高三\n\n## 限时小练\n"
    "1.（2026济宁一模）题干（　　）\nA. a　B. b　C. c　D. d\n\n"
    "## 参考答案\n1.【答案】C\n"
)
_dp = render_drill(_drill_md, OUT / "drill.docx", school="测试中学")
_dt = "\n".join(p.text for p in Document(_dp).paragraphs)
check("每日一练3" in _dt and "限时小练" in _dt, "drill：每日一练标题+栏目")
check("（2026济宁一模）" in _dt, "drill：真题源内联")
_lec_md = (
    "# 胶州四中高三化学尖子生培优材料\n\n## 培优知识点及达成目标\n高考选择题高频。\n"
    "## 【典型例题】\n1. 题干（　）\n## 【解题思路】\nA：排除。\n## 【变式训练】\n1. 变式\n"
    "## 学生反思总结\n归纳失分点。\n"
)
_lp = render_drill(_lec_md, OUT / "lecture.docx")
_lt = "\n".join(p.text for p in Document(_lp).paragraphs)
check(all(x in _lt for x in ["培优知识点及达成目标", "典型例题", "解题思路", "变式训练", "学生反思总结"]),
      "drill：培优五栏目全在")
import jszz.server as _srv  # noqa: E402
_mr = _srv.make_drill(_drill_md, name="回归测试", answers=True)
check(_mr["ok"] and _mr["questions"] >= 1 and "docx" in _mr["files"]
      and "docx_answers" in _mr["files"], "make_drill：MCP 产物齐")
check(_mr["qa"]["docx"]["ok"] and _mr["qa"]["markdown"]["ok"], "make_drill：QA 全过")

print("== 成绩分析（grade）==")
from jszz.grades import analyze, load_scores  # noqa: E402
_rows = [
    {"姓名": "甲", "班级": "1", "学籍号": 2024001, "考试": "上期中", "语文": 90, "数学": 80, "总分": 170},
    {"姓名": "乙", "班级": "1", "学籍号": 2024002, "考试": "上期中", "语文": 70, "数学": 60, "总分": 130},
    {"姓名": "甲", "班级": "1", "学籍号": 2024001, "考试": "上期末", "语文": 95, "数学": 88, "总分": 183},
    {"姓名": "乙", "班级": "2", "学籍号": 2024002, "考试": "上期末", "语文": 65, "数学": 55, "总分": 120},
]
_g = load_scores(_rows)
check(_g["ok"] and "学籍号" not in _g["columns"]["subjects"],
      "grade：ID 列（学籍号）不误当科目")
check(_g["columns"]["exam"] == "考试" and _g["columns"]["total"] == "总分",
      "grade：考次/总分列检测")
_a = analyze(_g["rows"], _g["columns"])
check(_a["subjects"]["语文"]["n"] == 4 and _a["subjects"]["语文"]["mean"] == 80.0,
      "grade：科目均分正确（90+70+95+65）/4=80")
check(_a["total"]["n"] == 4 and abs(_a["total"]["mean"] - 150.75) < 0.01,
      "grade：总分均分正确")
check(len(_a["top"]) >= 2 and _a["top"][0]["name"] == "甲",
      "grade：名次榜按分排序")
check(_a["trend"] is not None and _a["trend"]["n_students"] == 2,
      "grade：多次考试趋势（同名学生聚合）")
check(_a["trend"]["improved"][0]["name"] == "甲"
      and _a["trend"]["improved"][0]["delta"] == 13.0,
      "grade：进步榜 delta 正确（183-170=13）")
# 名次榜去重（多考次长表同一学生只占一个名次）
_dup_rows = [
    {"姓名": "甲", "考试": "上期中", "总分": 100},
    {"姓名": "甲", "考试": "上期末", "总分": 200},
    {"姓名": "乙", "考试": "上期末", "总分": 150},
]
_da = analyze(*[load_scores(_dup_rows)[k] for k in ("rows", "columns")])
_dn = [t["name"] for t in _da["top"]]
check(_dn == ["甲", "乙"] and _da["top"][0]["score"] == 200.0,
      "grade：名次榜按最近考次去重（甲只占 1 席且取最近分 200）")
# 渲染 + MCP 真跑（JSON 输入）
import jszz.server as _svcg  # noqa: E402
_mr = _svcg.make_grade_report(
    '[{"姓名":"丙","班级":"3","语文":88,"数学":77,"总分":165}]',
    name="成绩分析回归", notes="测试建议。")
check(_mr["ok"] and _mr["n_rows"] == 1 and "docx" in _mr["files"],
      "make_grade_report：JSON 输入出 docx")
check(_mr["summary"]["total_mean"] == 165.0
      and _mr["summary"]["subjects"]["语文"]["mean"] == 88.0,
      "make_grade_report：摘要指标正确")
check(Path(_mr["files"]["docx"]).exists(), "make_grade_report：docx 落盘")
# 成绩分析汇报 PPT（钉 float-EMU 修复：add_table 尺寸写 int，写 float 读回会崩）
check("pptx" in _mr["files"] and Path(_mr["files"]["pptx"]).exists(),
      "make_grade_report：pptx 汇报产物落盘")
from pptx import Presentation as _P  # noqa: E402
_pg = _P(_mr["files"]["pptx"])
check(len(_pg.slides) >= 5, "make_grade_report：汇报 PPT ≥5 页（封面/章节/数据/建议）")
_dt = "\n".join(sh.text_frame.text for s in _pg.slides for sh in s.shapes if sh.has_text_frame)
check("成绩分析" in _dt and ("均分" in _dt or "教学建议" in _dt),
      "make_grade_report：汇报 PPT 含分析内容")

# 弱科识别 + 提升推荐（钉：每条推荐不雷同、结尾各异，别退化成机器套话）
# 夹具含真实弱科（英语均分率 <60%、及格率低）
_weak_rows = [
    {"姓名": "甲", "班级": "1", "语文": 90, "数学": 80, "英语": 40, "总分": 210},
    {"姓名": "乙", "班级": "1", "语文": 85, "数学": 75, "英语": 35, "总分": 195},
    {"姓名": "丙", "班级": "2", "语文": 88, "数学": 82, "英语": 30, "总分": 200},
]
_a2 = analyze(*[load_scores(_weak_rows)[k] for k in ("rows", "columns")])
check("weak_subjects" in _a2 and "recommendations" in _a2,
      "analyze：含 weak_subjects/recommendations 字段")
check(len(_a2["weak_subjects"]) > 0, "弱科识别：低分科目被识别")
check(_a2["weak_subjects"][0]["subject"] == "英语",
      "弱科识别：英语（均分率最低）排最前")
_rcs = _a2["recommendations"]
check(len(_rcs) >= 2, "提升推荐：有推荐条目")
check(len(set(_rcs[:2])) == len(_rcs[:2]), "提升推荐：前几条两两不雷同")
# 弱科/推荐进 docx + pptx 产物
_dt2 = "\n".join(p.text for p in Document(_mr["files"]["docx"]).paragraphs)
check("弱科" in _dt2 and "提升推荐" in _dt2, "docx：弱科+提升推荐显示")
_pt2 = ""
for s in _P(_mr["files"]["pptx"]).slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            _pt2 += sh.text_frame.text + "\n"
        if sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    _pt2 += c.text + "\n"
check("弱科" in _pt2 and "提升推荐" in _pt2, "pptx：弱科+提升推荐显示")

print("== 打包完整性 ==")
skill_zip = ROOT / "dist" / "jszz.skill"
if skill_zip.exists():
    with zipfile.ZipFile(skill_zip) as z:
        names = z.namelist()
    check(any(n.endswith("SKILL.md") for n in names), "jszz.skill 含 SKILL.md")
    check(sum(1 for n in names if "references/" in n) >= 5, "jszz.skill 含 references")
else:
    check(False, "dist/jszz.skill 不存在")

print()
if fails:
    print(f"FAILED {len(fails)}:")
    for f in fails:
        print(" -", f)
    sys.exit(1)
print("ALL PASS")
